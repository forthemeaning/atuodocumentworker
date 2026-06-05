#!/usr/bin/env python3
"""Batch import images into Word table with OCR text extraction."""

import os
# Workaround for OpenMP conflict between PyTorch and other DLLs
os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"

import argparse
import sys
from PIL import Image
from docx import Document
from docx.shared import Cm, Emu
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement

try:
    import easyocr
    HAS_EASYOCR = True
except ImportError:
    HAS_EASYOCR = False


SUPPORTED_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff"}


def parse_args():
    parser = argparse.ArgumentParser(description="Batch import images into Word table with OCR")
    parser.add_argument("--folder", required=True, help="Source folder containing images")
    parser.add_argument("--output", default="output.docx", help="Output Word file path")
    parser.add_argument("--width", type=float, default=8, help="Total table width in cm (default: 8)")
    parser.add_argument("--height", type=float, default=3, help="Row height in cm (default: 3)")
    parser.add_argument("--no-ocr", action="store_true", help="Skip OCR (image only)")
    return parser.parse_args()


def get_image_files(folder):
    if not os.path.isdir(folder):
        print(f"Error: folder '{folder}' does not exist.")
        sys.exit(1)
    files = sorted(
        f for f in os.listdir(folder)
        if os.path.splitext(f)[1].lower() in SUPPORTED_EXT
    )
    if not files:
        print(f"Error: no supported image files found in '{folder}'.")
        sys.exit(1)
    return files


def calc_scaled_size(img_path, max_w_emu, max_h_emu):
    with Image.open(img_path) as img:
        px_w, px_h = img.size
    native_w_emu = int(px_w * 914400 / 72)
    native_h_emu = int(px_h * 914400 / 72)
    scale = min(max_w_emu / native_w_emu, max_h_emu / native_h_emu)
    return int(native_w_emu * scale), int(native_h_emu * scale)


def _set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "single")
        element.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz", "4")
        element.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space", "0")
        element.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color", "000000")
        borders.append(element)
    tblPr.append(borders)


def _setup_cell(cell, is_text=False):
    """Remove cell margins and paragraph spacing."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    # Remove existing tcMar if any
    old = tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcMar")
    if old is not None:
        tcPr.remove(old)
    # Set zero margins
    tcMar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w", str(is_text * 57))  # ~0.1cm if text cell
        el.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type", "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)
    # Zero paragraph spacing
    pPr = cell.paragraphs[0]._p.get_or_add_pPr()
    old_sp = pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing")
    if old_sp is not None:
        pPr.remove(old_sp)
    spacing = OxmlElement("w:spacing")
    spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}before", "0")
    spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}after", "0")
    spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line", "240")
    spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lineRule", "auto")
    pPr.append(spacing)


def insert_image_cell(cell, filepath, max_w_emu, max_h_emu, margin=0.9):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _setup_cell(cell)
    p = cell.paragraphs[0]
    # Remove any existing content but keep paragraph
    for r in p.runs:
        r._r.getparent().remove(r._r)
    try:
        img_w_emu, img_h_emu = calc_scaled_size(
            filepath, int(max_w_emu * margin), int(max_h_emu * margin)
        )
    except Exception:
        p.text = "[Error]"
        return False
    run = p.add_run()
    run.add_picture(filepath, width=Emu(img_w_emu), height=Emu(img_h_emu))
    return True


def insert_text_cell(cell, text):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _setup_cell(cell, is_text=True)
    p = cell.paragraphs[0]
    for r in p.runs:
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    run.font.size = Cm(0.3)


def ocr_image(reader, filepath):
    """Extract text from image. Returns (success_bool, text)."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".gif":
        # EasyOCR doesn't support GIF; try converting first frame
        try:
            from PIL import Image
            gif = Image.open(filepath)
            frame = gif.copy().convert("RGB")
            import tempfile
            tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
            tmp_path = tmp.name
            tmp.close()
            frame.save(tmp_path)
            result = reader.readtext(tmp_path)
            os.unlink(tmp_path)
        except Exception as e:
            return False, str(e)
    else:
        try:
            result = reader.readtext(filepath)
        except Exception as e:
            return False, str(e)
    lines = []
    for _, text, confidence in result:
        if confidence >= 0.3:
            lines.append(text.strip())
    text = "\n".join(lines) if lines else "(no text detected)"
    return True, text


def main(folder=None, output=None, width=None, height=None, no_ocr=False):
    if folder is not None:
        folder_path = folder
        output_path = output or "output.docx"
        cell_w = width or 8
        cell_h = height or 3
        skip_ocr = no_ocr
    else:
        args = parse_args()
        folder_path = args.folder
        output_path = args.output
        cell_w = args.width
        cell_h = args.height
        skip_ocr = args.no_ocr

    total_width_emu = Cm(cell_w).emu
    row_height_emu = Cm(cell_h).emu
    col_width_emu = total_width_emu // 2

    image_files = get_image_files(folder_path)

    # Init OCR engine
    reader = None
    if not skip_ocr:
        if not HAS_EASYOCR:
            print("Warning: easyocr not installed. Install with: pip install easyocr")
            print("Falling back to image-only mode.\n")
            skip_ocr = True
        else:
            print("Loading OCR engine (first run downloads model)...")
            reader = easyocr.Reader(["ch_sim", "en"], gpu=False)
            print("OCR ready.\n")

    doc = Document()

    table = doc.add_table(rows=len(image_files), cols=2)
    table.autofit = False
    _set_table_borders(table)
    table.columns[0].width = Emu(col_width_emu)
    table.columns[1].width = Emu(col_width_emu)

    for i, filename in enumerate(image_files):
        filepath = os.path.join(folder_path, filename)
        row = table.rows[i]
        row.height = Emu(row_height_emu)

        img_cell = row.cells[0]
        img_cell.width = Emu(col_width_emu)
        insert_image_cell(img_cell, filepath, col_width_emu, row_height_emu)

        # OCR
        if not skip_ocr and reader is not None:
            print(f"  {filename} (OCR...", end="", flush=True)
            ok, text = ocr_image(reader, filepath)
            print(f" {len(text)} chars)")
            txt_cell = row.cells[1]
            txt_cell.width = Emu(col_width_emu)
            insert_text_cell(txt_cell, text)
        else:
            row.cells[1].text = ""
            print(f"  {filename}")

    script_dir = os.path.dirname(os.path.abspath(__file__))
    output_file = os.path.join(script_dir, output_path) if not os.path.isabs(output_path) else output_path
    doc.save(output_file)
    print(f"\nDone! {len(image_files)} images -> '{output_file}'")


if __name__ == "__main__":
    main(folder=r"C:\Users\YxYxi\Pictures\Everything is Crab\Run #3")
